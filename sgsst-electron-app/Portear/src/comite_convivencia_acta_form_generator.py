import json
import sys
import os
from openpyxl import load_workbook
from datetime import datetime

def log_message(message, level="INFO"):
    """Generates a JSON-formatted log message."""
    print(json.dumps({"type": "log", "message": message, "level": level}))

def main(temp_data_path, output_path):
    """
    Loads an Excel template for Comité de Convivencia, applies changes from a JSON file
    with structured data from the form interface, and saves the result to a new file,
    preserving styles.
    """
    try:
        log_message("Iniciando generación de acta de Comité de Convivencia desde formulario...")

        # 1. Cargar datos del JSON temporal
        with open(temp_data_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        # Extraer los datos estructurados
        acta_number = data.get('actaNumber', '000')
        fecha = data.get('fecha', '')
        inicia = data.get('inicia', '00:00')
        termina = data.get('termina', '00:00')
        topic = data.get('topic', '')
        lugar = data.get('lugar', '')
        agenda = data.get('agenda', [])
        desarrollo = data.get('desarrollo', [])
        participantes = data.get('participantes', [])
        conclusiones = data.get('conclusiones', [])
        compromisos = data.get('compromisos', [])
        
        log_message(f"Datos extraídos: acta_number={acta_number}, fecha={fecha}, topic={topic}")

        # 2. Ruta a la plantilla
        template_path = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), 'utils', 'GI-FO-029 ACTA DE REUNION CONVIVENCIA Mayo.xlsx')
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Plantilla no encontrada en: {template_path}")

        log_message(f"Cargando plantilla desde: {template_path}")

        # 3. Cargar la plantilla con openpyxl
        wb = load_workbook(template_path)
        ws = wb.active

        # 4. Aplicar datos generales del acta
        # Suponiendo que la plantilla tiene celdas específicas donde se colocan estos datos
        # Ajustar las celdas según la estructura real de la plantilla
        try:
            ws['B8'] = f"N° {acta_number}"  # Número de acta
        except:
            log_message("No se pudo encontrar la celda para el número de acta (B8)", "WARN")
        
        try:
            ws['B9'] = fecha  # Fecha de la reunión
        except:
            log_message("No se pudo encontrar la celda para la fecha (B9)", "WARN")
            
        try:
            ws['B10'] = f"{inicia} - {termina}"  # Horario
        except:
            log_message("No se pudo encontrar la celda para el horario (B10)", "WARN")
            
        try:
            ws['B11'] = topic  # Tema de la reunión
        except:
            log_message("No se pudo encontrar la celda para el tema (B11)", "WARN")
            
        try:
            ws['B12'] = lugar  # Lugar de la reunión
        except:
            log_message("No se pudo encontrar la celda para el lugar (B12)", "WARN")

        # 5. Aplicar agenda items
        start_row_agenda = 17  # Suponiendo que la agenda comienza en la fila 17
        for i, item in enumerate(agenda):
            try:
                current_row = start_row_agenda + i
                ws.cell(row=current_row, column=2, value=item.get('tema', ''))  # Columna B: Tema
                ws.cell(row=current_row, column=3, value=item.get('responsable', ''))  # Columna C: Responsable
                ws.cell(row=current_row, column=4, value=item.get('tiempo', ''))  # Columna D: Tiempo
            except Exception as e:
                log_message(f"Error al aplicar agenda item {i}: {str(e)}", "WARN")

        # 6. Aplicar desarrollo items
        start_row_desarrollo = 25  # Suponiendo que el desarrollo comienza en la fila 25
        for i, item in enumerate(desarrollo):
            try:
                current_row = start_row_desarrollo + i
                ws.cell(row=current_row, column=2, value=item.get('asunto', ''))  # Columna B: Asunto
                ws.cell(row=current_row, column=3, value=item.get('desarrollo', ''))  # Columna C: Desarrollo
            except Exception as e:
                log_message(f"Error al aplicar desarrollo item {i}: {str(e)}", "WARN")

        # 7. Aplicar participantes
        start_row_participantes = 30  # Suponiendo que los participantes comienzan en la fila 30
        for i, participante in enumerate(participantes):
            try:
                current_row = start_row_participantes + i
                ws.cell(row=current_row, column=2, value=participante)  # Columna B: Participante
            except Exception as e:
                log_message(f"Error al aplicar participante {i}: {str(e)}", "WARN")

        # 8. Aplicar conclusiones
        start_row_conclusiones = 40  # Suponiendo que las conclusiones comienzan en la fila 40
        for i, conclusion in enumerate(conclusiones):
            try:
                current_row = start_row_conclusiones + i
                ws.cell(row=current_row, column=2, value=conclusion)  # Columna B: Conclusión
            except Exception as e:
                log_message(f"Error al aplicar conclusión {i}: {str(e)}", "WARN")

        # 9. Aplicar compromisos
        start_row_compromisos = 45  # Suponiendo que los compromisos comienzan en la fila 45
        for i, compromiso in enumerate(compromisos):
            try:
                current_row = start_row_compromisos + i
                ws.cell(row=current_row, column=2, value=compromiso.get('descripcion', ''))  # Columna B: Descripción
                ws.cell(row=current_row, column=3, value=compromiso.get('responsable', ''))  # Columna C: Responsable
                ws.cell(row=current_row, column=4, value=compromiso.get('fecha', ''))  # Columna D: Fecha
            except Exception as e:
                log_message(f"Error al aplicar compromiso {i}: {str(e)}", "WARN")

        log_message("Todos los datos han sido aplicados a la plantilla en memoria.")

        # 10. Guardar el nuevo archivo en la ruta de salida especificada por el usuario
        wb.save(output_path)
        log_message(f"Acta de Comité de Convivencia generada exitosamente en: {output_path}")

        # 11. Devolver resultado de éxito
        print(json.dumps({
            "type": "result",
            "payload": {
                "success": True,
                "documentPath": output_path
            }
        }))

    except Exception as e:
        log_message(f"Error crítico generando el acta: {str(e)}", "ERROR")
        print(json.dumps({
            "type": "result",
            "payload": {
                "success": False,
                "error": str(e)
            }
        }))
        sys.exit(1)

if __name__ == "__main__":
    # El script espera 2 argumentos: la ruta al JSON temporal y la ruta de salida final
    if len(sys.argv) < 3:
        log_message("Uso: python comite_convivencia_acta_form_generator.py <ruta_json_temporal> <ruta_salida_excel>", "ERROR")
        sys.exit(1)

    temp_json_path = sys.argv[1]
    final_output_path = sys.argv[2]
    main(temp_json_path, final_output_path)